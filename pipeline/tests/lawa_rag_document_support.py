"""Small in-memory Office fixtures and a parser process that simulates a hang."""

import time
from io import BytesIO


def mixed_documents():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    document = Document()
    document.add_heading("Admissions", 1)
    document.add_paragraph("شروط القبول: تقديم الطلب عبر الموقع.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Program", "Years"), ("Robotics", "2"))):
        for cell, text in zip(row.cells, values):
            cell.text = text
    word = BytesIO()
    document.save(word)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Scholarships"
    slide.placeholders[1].text = "المنح الدراسية متاحة للطلاب المؤهلين."
    presentation = BytesIO()
    slides.save(presentation)
    book = Workbook()
    book.active.title = "Programs"
    book.active.append(["Code", "Program"])
    book.active.append(["007", "Artificial intelligence"])
    spreadsheet = BytesIO()
    book.save(spreadsheet)
    return {
        "admissions.docx": word.getvalue(), "scholarships.pptx": presentation.getvalue(),
        "programs.xlsx": spreadsheet.getvalue(),
        "codes.csv": 'code,name\n007,الذكاء الاصطناعي\n'.encode(),
        "guide.html": b"<html><body><h1>Study</h1><p>Apply online.</p></body></html>",
    }


def hang_worker(connection, _):
    connection.recv()
    time.sleep(30)


def pdf_bytes(*page_sizes):
    import pypdfium2 as pdfium

    output = BytesIO()
    with pdfium.PdfDocument.new() as document:
        for width, height in page_sizes:
            document.new_page(width, height).close()
        document.save(output)
    return output.getvalue()


def graceful_worker(connection, options):
    from pathlib import Path
    from lawa_rag.ingestion._document_worker import _convert_request
    from lawa_rag.ingestion import ConversionOptions

    suffix, source, target = connection.recv()
    _convert_request(suffix, Path(source), Path(target), ConversionOptions(**options))
    connection.send("complete")
    if connection.recv() is None:
        Path(options["docling_artifacts_path"]).joinpath("clean-shutdown").touch()
    connection.close()
