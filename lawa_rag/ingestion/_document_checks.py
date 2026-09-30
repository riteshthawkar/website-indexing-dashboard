"""Input budgets and completeness guards around upstream parsers, not new converters."""

from __future__ import annotations

import csv
import io
import math
import zipfile
from contextlib import ExitStack
from pathlib import Path, PurePosixPath

from .documents import ConversionFailure


SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".json", ".html", ".htm"}


def preflight(suffix, body, options):
    if suffix not in SUPPORTED:
        raise ConversionFailure("unsupported_format")
    if not body:
        raise ConversionFailure("empty_document")
    if suffix in {".docx", ".pptx", ".xlsx"}:
        _check_archive(body, options)
    if suffix == ".xlsx":
        _check_workbook(body, options)
    if suffix == ".csv":
        return _check_csv(body, options)
    if suffix == ".pdf":
        for value in (options.docling_artifacts_path, options.ocr_model_path):
            if not value or not Path(value).is_dir():
                raise ConversionFailure("ocr_models_not_provisioned")
        _check_pdf(body, options)
    return body


def _check_pdf(body, options):
    import pypdfium2 as pdfium

    try:
        with pdfium.PdfDocument(body) as document:
            if not 0 < len(document) <= options.max_pages:
                raise ConversionFailure("page_limit_exceeded")
            for index in range(len(document)):
                width, height = document.get_page_size(index)
                if not all(math.isfinite(value) and value > 0 for value in (width, height)):
                    raise ConversionFailure("invalid_pdf_page_geometry")
                # Pinned Docling EasyOCR renders at 3x (216 DPI), TableFormer at 2x.
                pixels = math.ceil(width * 3) * math.ceil(height * 3)
                if pixels > options.max_pdf_page_pixels:
                    raise ConversionFailure("pdf_page_pixel_limit_exceeded")
    except pdfium.PdfiumError:
        raise ConversionFailure("invalid_or_encrypted_pdf") from None


def _check_archive(body, options):
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000:
            raise ConversionFailure("archive_limit_exceeded")
        if sum(item.file_size for item in entries) > options.max_archive_bytes:
            raise ConversionFailure("archive_limit_exceeded")
        for item in entries:
            path = PurePosixPath(item.filename)
            if path.is_absolute() or ".." in path.parts or item.flag_bits & 1:
                raise ConversionFailure("unsupported_archive_entry")


def _check_csv(body, options):
    source = io.StringIO(body.decode("utf-8-sig", errors="strict"))
    output = io.StringIO()
    writer, width, cells = csv.writer(output), None, 0
    reader = csv.reader(source, delimiter=options.csv_delimiter, strict=True)
    for index, row in enumerate(reader, 1):
        width = len(row) if width is None else width
        if not width or len(row) != width:
            raise ConversionFailure("inconsistent_csv_columns")
        cells += len(row)
        if index > options.max_rows or cells > options.max_cells:
            raise ConversionFailure("table_limit_exceeded")
        writer.writerow(row)
    return output.getvalue().encode("utf-8")


def _check_workbook(body, options):
    from openpyxl import load_workbook

    with ExitStack() as stack:
        books = []
        for mode in (False, True):
            book = load_workbook(io.BytesIO(body), read_only=True, data_only=mode, keep_links=False)
            stack.callback(book.close)
            books.append(book)
        rows, cells = 0, 0
        for sheet, values in zip(books[0].worksheets, books[1].worksheets):
            rows += sheet.max_row or 0
            cells += (sheet.max_row or 0) * (sheet.max_column or 0)
            if rows > options.max_rows or cells > options.max_cells:
                raise ConversionFailure("table_limit_exceeded")
            for row, cached in zip(sheet.iter_rows(), values.iter_rows()):
                for cell, value in zip(row, cached):
                    if cell.data_type == "f" and value.value is None:
                        raise ConversionFailure("formula_cache_missing")
